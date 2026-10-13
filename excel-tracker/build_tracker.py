#!/usr/bin/env python3
"""Генератор трекера арбітражу (Nutra / CPA) у форматі xlsx.

Запуск:  python3 build_tracker.py [папка_для_файлів]
Створює два файли:
  arbitrage_tracker_demo.xlsx — з демо-даними, щоб побачити, як усе працює;
  arbitrage_tracker.xlsx      — чистий шаблон (по одному рядку-прикладу).
"""
import datetime as dt
import math
import random
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.comments import Comment
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import CharacterProperties, Font as DFont, Paragraph, ParagraphProperties
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink

# ───────────────────────── палітра (темна тема) ─────────────────────────
FONT, MONO = "Arial", "Consolas"
BG, PANEL, HEAD, BORDER = "0E1220", "161B2E", "1F2640", "2A3252"
INPUT, INPUT_TXT = "16243F", "7DB3FF"
TXT, MUTED, DIM, ACCENT, HEAD_TXT = "E6E9F2", "8A93B0", "4A5272", "8B9CFF", "C3CAE0"
GREEN, GREEN_BG = "34D399", "0F2E25"
YELLOW, YELLOW_BG = "FBBF24", "3A2E0E"
RED, RED_BG, DRED_BG = "F87171", "3D171B", "991B1B"
GRAY, GRAY_BG = "A5B4FC", "232A4A"
BAR1, BAR2 = "5B6CFF", "34D399"

# ───────────────────────── формати чисел ─────────────────────────
USD = '$#,##0.00;-$#,##0.00;"–"'
USD0 = '$#,##0;-$#,##0;"–"'
PROFIT = '+$#,##0.00;-$#,##0.00;"–"'
PROFIT0 = '+$#,##0;-$#,##0;"–"'
INT = '#,##0;-#,##0;"–"'
INT_IN, USD_IN = "#,##0", "$#,##0.00"   # у синіх клітинках нуль показуємо як 0
DEC1 = '#,##0.0;-#,##0.0;"–"'
PCT0 = '0%;-0%;"–"'
PCT1 = '0.0%;-0.0%;"–"'
ROI = '+0%;-0%;0%'
DATE = "dd.mm.yyyy"
MULT = '0.0" × Max CPL"'

# ───────────────────────── аркуші ─────────────────────────
S_START, S_CALC, S_STAT, S_LOG = "Старт", "Гео-калькулятор", "Статистика", "Журнал дій"
S_DAY, S_GEO, S_KILL, S_CRE, S_LIST = (
    "Загальна по днях", "Зведення по гео", "Спенд без лідів", "Креативи", "Списки")


def ref(sheet, rng):
    return f"'{sheet}'!{rng}"


N_GEO, GEO_R0 = 7, 7                      # гео в калькуляторі: рядки 7..13
GEO_R1 = GEO_R0 + N_GEO - 1
STAT_R0, STAT_R1 = 7, 1006                # 1000 рядків статистики
LOG_R0, LOG_R1 = 9, 308                   # 300 рядків журналу
CRE_R0, CRE_R1 = 8, 207                   # 200 креативів
DAY_R0, DAY_N = 11, 31                    # 31 день у загальній таблиці
DAY_R1 = DAY_R0 + DAY_N - 1
CHK_R0, CHK_R1 = 25, 44                   # перевірка адсетів


def calc_col(col):
    return ref(S_CALC, f"${col}${GEO_R0}:${col}${GEO_R1}")


def stat_col(col):
    return ref(S_STAT, f"${col}${STAT_R0}:${col}${STAT_R1}")


GEO_LIST = calc_col("B")
T1, T2, T3 = ref(S_KILL, "$H$5"), ref(S_KILL, "$H$6"), ref(S_KILL, "$H$7")


def lookup(col, key):
    """Значення з калькулятора по коду гео."""
    return f"INDEX({calc_col(col)},MATCH({key},{GEO_LIST},0))"


V_EARLY, V_Q, V_OFF, V_LATE = "Рано судити", "Під питанням", "Вимикати", "Давно треба було"
V_OK, V_SCALE = "В цілі", "Масштабувати"


def ladder(p):
    """Вердикт за шансом, що адсет нормальний (пороги — на аркуші «Спенд без лідів»)."""
    return (f'IF({p}>={T1},"{V_EARLY}",IF({p}>={T2},"{V_Q}",'
            f'IF({p}>={T3},"{V_OFF}","{V_LATE}")))')


# ───────────────────────── стилі ─────────────────────────
def fill(c):
    return PatternFill("solid", start_color=c, end_color=c)


def font(color=TXT, size=10, bold=False, italic=False, name=FONT, underline=None):
    return Font(name=name, size=size, bold=bold, italic=italic, color=color, underline=underline)


def side(c=BORDER, style="thin"):
    return Side(style=style, color=c)


ROW_LINE = Border(bottom=side())


def paint(ws, rows, cols=30):
    f, fn = fill(BG), font()
    for r in range(1, rows + 1):
        for c in range(1, cols + 1):
            cell = ws.cell(r, c)
            cell.fill, cell.font = f, fn
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2.5
    ws.row_dimensions[1].height = 8
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def widths(ws, mapping):
    for col, w in mapping.items():
        ws.column_dimensions[col].width = w


def title(ws, text, sub):
    ws["B2"], ws["B3"] = text, sub
    ws["B2"].font = font(TXT, 18, True)
    ws["B3"].font = font(MUTED, 10)
    ws.row_dimensions[2].height = 28
    ws.row_dimensions[3].height = 16


def group(ws, row, c1, c2, text, color=ACCENT):
    if c2 > c1:
        ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
    c = ws.cell(row, c1, text)
    c.font, c.fill = font(color, 8, True), fill(PANEL)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for col in range(c1, c2 + 1):
        ws.cell(row, col).fill = fill(PANEL)
        ws.cell(row, col).border = Border(top=side(color, "medium"))
    ws.row_dimensions[row].height = 18


def headers(ws, row, c1, names, height=36):
    for i, h in enumerate(names):
        c = ws.cell(row, c1 + i, h)
        c.fill, c.font = fill(HEAD), font(HEAD_TXT, 9, True)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(bottom=side(ACCENT, "medium"))
    ws.row_dimensions[row].height = height


def inp(c, fmt=None, align="right"):
    c.fill, c.font, c.border = fill(INPUT), font(INPUT_TXT), ROW_LINE
    c.alignment = Alignment(horizontal=align, vertical="center")
    if fmt:
        c.number_format = fmt


def calc(c, fmt=None, align="right", bold=False, color=TXT):
    c.fill, c.font, c.border = fill(PANEL), font(color, bold=bold), ROW_LINE
    c.alignment = Alignment(horizontal=align, vertical="center")
    if fmt:
        c.number_format = fmt


def label(c, text, color=MUTED, size=9, bold=False, align="left", wrap=False):
    c.value = text
    c.font = font(color, size, bold)
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)


def note(cell, text):
    cm = Comment(text, "Трекер")
    cm.width, cm.height = 260, 110
    cell.comment = cm


def merge(ws, rng):
    ws.merge_cells(rng)
    return ws[rng.split(":")[0]]


def cf(ws, rng, formula, color, bg=None, bold=True):
    kw = {"font": Font(color=color, bold=bold)}
    if bg:
        kw["fill"] = fill(bg)
    ws.conditional_formatting.add(rng, FormulaRule(formula=[formula], stopIfTrue=True, **kw))


VERDICT_STYLE = [
    (V_SCALE, BG, GREEN), (V_OK, GREEN, GREEN_BG), (V_EARLY, GRAY, GRAY_BG),
    (V_Q, YELLOW, YELLOW_BG), (V_OFF, RED, RED_BG), (V_LATE, "FFFFFF", DRED_BG),
]


def cf_values(ws, rng, anchor, styles):
    for val, color, bg in styles:
        cf(ws, rng, f'{anchor}="{val}"', color, bg)


def cf_sign(ws, rng, anchor):
    cf(ws, rng, f"AND(ISNUMBER({anchor}),{anchor}<0)", RED, bold=False)
    cf(ws, rng, f"AND(ISNUMBER({anchor}),{anchor}>0)", GREEN, bold=False)


def cf_cpl_ratio(ws, rng, anchor):
    cf(ws, rng, f"AND(ISNUMBER({anchor}),{anchor}<=1)", GREEN, GREEN_BG)
    cf(ws, rng, f"AND(ISNUMBER({anchor}),{anchor}<=1.3)", YELLOW, YELLOW_BG)
    cf(ws, rng, f"ISNUMBER({anchor})", RED, RED_BG)


def dv_list(ws, rng, source, strict=False):
    dv = DataValidation(type="list", formula1=f"={source}", allow_blank=True)
    dv.errorStyle = "stop" if strict else "warning"
    dv.errorTitle = "Немає в списку"
    dv.error = ("Такого гео немає в «Гео-калькулятор». Додайте його туди, інакше Max CPL не підтягнеться."
                if source == GEO_LIST else "Оберіть значення зі списку (список редагується на аркуші «Списки»).")
    dv.showErrorMessage = True
    ws.add_data_validation(dv)
    dv.add(rng)


def dv_range(ws, rng, lo, hi, msg):
    dv = DataValidation(type="decimal", operator="between", formula1=str(lo), formula2=str(hi), allow_blank=True)
    dv.errorTitle, dv.error, dv.showErrorMessage = "Перевірте значення", msg, True
    ws.add_data_validation(dv)
    dv.add(rng)


# ───────────────────────── графіки ─────────────────────────
def rich(color, size=900, bold=False):
    cp = CharacterProperties(sz=size, b=bold, solidFill=color, latin=DFont(typeface=FONT))
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def dark_chart(ch, axes):
    ch.graphical_properties = GraphicalProperties(solidFill=PANEL, ln=LineProperties(solidFill=BORDER))
    ch.plot_area.graphicalProperties = GraphicalProperties(noFill=True, ln=LineProperties(noFill=True))
    ch.legend.position = "t"
    ch.legend.txPr = rich(TXT, 900)
    for ax in axes:
        ax.delete = False
        ax.txPr = rich(MUTED, 800)
        ax.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=BORDER))
        if ax.majorGridlines is not None:
            ax.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=BORDER, prstDash="dash"))


def bar_colors(ch, colors):
    for s, c in zip(ch.series, colors):
        s.graphicalProperties = GraphicalProperties(solidFill=c, ln=LineProperties(noFill=True))


# ═════════════════════════ АРКУШІ ═════════════════════════
def build_calc(ws, geos):
    paint(ws, 40)
    title(ws, "Гео-калькулятор",
          "Сині клітинки — дані по офферу. Все праворуч рахується саме і використовується на всіх інших аркушах.")
    widths(ws, {"B": 8, "C": 26, "D": 12, "E": 11, "F": 11, "G": 11, "H": 10, "I": 11, "J": 12,
                "K": 12, "L": 12, "M": 11, "N": 11, "O": 13, "P": 14, "Q": 30})
    group(ws, 5, 2, 8, "ВХІДНІ ДАНІ — ЗАПОВНЮЄТЕ ВИ")
    group(ws, 5, 9, 16, "РОЗРАХУНОК — АВТОМАТИЧНО", GREEN)
    group(ws, 5, 17, 17, "")
    headers(ws, 6, 2, ["Гео", "Країна / оффер", "Виплата за апрув, $", "% апруву (план)", "Цільовий ROI",
                       "CR клік → лід", "CTR", "EPL, $", "CPL беззбитку, $", "Max CPL, $",
                       "Max ціна апруву, $", "Max CPC, $", "Max CPM, $", "Не судити до, $",
                       "Вимикати після\n(0 лідів), $", "Нотатки"], 44)
    for addr, text in {
        "D6": "Скільки партнерка платить за 1 апрувнутий (підтверджений) лід.",
        "E6": "Очікуваний % апруву. Фактичний апрув видно на аркуші «Зведення по гео».",
        "F6": "Скільки хочете заробити понад спенд. 30% = на кожні $100 спенду $30 чистого профіту.",
        "G6": "Яка частка кліків по посиланню стає лідами (конверсія ленда/преленда).",
        "H6": "CTR по посиланню (link CTR) з рекламного кабінету.",
        "I6": "Earnings per lead — скільки в середньому приносить 1 лід = Виплата × % апруву.",
        "K6": "Головна цифра: ліміт CPL, при якому виходить цільовий ROI = EPL ÷ (1 + ROI).",
        "O6": "Поки адсет відкрутив без лідів менше цієї суми — шанс, що він нормальний, ще вище порогу "
              "«Рано судити». Пороги — на аркуші «Спенд без лідів».",
        "P6": "Якщо адсет відкрутив стільки і не дав жодного ліда — шанс, що він нормальний, нижче порогу "
              "«Під питанням». Вимикайте.",
    }.items():
        note(ws[addr], text)

    for i in range(N_GEO):
        r = GEO_R0 + i
        g = geos[i] if i < len(geos) else (None,) * 8
        fmts = [None, None, USD_IN, "0%", "0%", "0.0%", "0.0%"]
        for j, (val, fmt) in enumerate(zip(g[:7], fmts)):
            c = ws.cell(r, 2 + j, val)
            inp(c, fmt, "center" if j == 0 else ("left" if j == 1 else "right"))
        c = ws.cell(r, 17, g[7])
        inp(c, None, "left")
        f = {
            "I": f'=IF(OR($D{r}="",$E{r}=""),"",$D{r}*$E{r})',
            "J": f'=IF(I{r}="","",I{r})',
            "K": f'=IF(I{r}="","",I{r}/(1+N($F{r})))',
            "L": f'=IF($D{r}="","",$D{r}/(1+N($F{r})))',
            "M": f'=IF(OR(K{r}="",$G{r}=""),"",K{r}*$G{r})',
            "N": f'=IF(OR(M{r}="",$H{r}=""),"",M{r}*$H{r}*1000)',
            "O": f'=IF(K{r}="","",-LN({T1})*K{r})',
            "P": f'=IF(K{r}="","",-LN({T2})*K{r})',
        }
        for col, formula in f.items():
            ws[f"{col}{r}"] = formula
            calc(ws[f"{col}{r}"], USD, bold=(col == "K"), color=(GREEN if col == "K" else TXT))
        ws.row_dimensions[r].height = 20

    rows = GEO_R0, GEO_R1
    dv_range(ws, f"E{rows[0]}:E{rows[1]}", 0, 1, "Введіть відсоток, напр. 40%.")
    dv_range(ws, f"G{rows[0]}:H{rows[1]}", 0, 1, "Введіть відсоток, напр. 5%.")
    dv_range(ws, f"F{rows[0]}:F{rows[1]}", -0.99, 20, "Введіть відсоток, напр. 30%.")

    r = GEO_R1 + 3
    label(ws.cell(r, 2), "ЯК РАХУЄТЬСЯ", ACCENT, 9, True)
    expl = [
        ("EPL", "Виплата × % апруву", "Скільки в середньому заробляєте з 1 ліда."),
        ("CPL беззбитку", "дорівнює EPL", "Якщо CPL вищий — працюєте в мінус."),
        ("Max CPL", "EPL ÷ (1 + цільовий ROI)", "Ліміт CPL, при якому виходить ваш цільовий ROI."),
        ("Max ціна апруву", "Виплата ÷ (1 + цільовий ROI)", "Скільки можна витратити на 1 апрув."),
        ("Max CPC", "Max CPL × CR", "Максимальна ціна кліка при вашій конверсії ленда."),
        ("Max CPM", "Max CPC × CTR × 1000", "Максимальна ціна 1000 показів при вашому CTR."),
        ("Не судити до", "−ln(поріг «Рано судити») × Max CPL", "При 30% ≈ 1.2 × Max CPL."),
        ("Вимикати після", "−ln(поріг «Під питанням») × Max CPL", "При 10% ≈ 2.3 × Max CPL без жодного ліда."),
    ]
    for k, (name, formula, text) in enumerate(expl):
        rr = r + 1 + k
        label(merge(ws, f"B{rr}:C{rr}"), name, TXT, 9, True)
        c = merge(ws, f"D{rr}:H{rr}")
        c.value, c.font = formula, font(ACCENT, 9, name=MONO)
        label(merge(ws, f"I{rr}:Q{rr}"), text)
    ws.freeze_panes = "D7"


def build_kill(ws, checks):
    paint(ws, 50)
    title(ws, "Спенд без лідів: вимикати чи чекати?",
          "Скільки адсет може відкрутити без лідів, поки це ще нормальна випадковість, а не поганий адсет.")
    widths(ws, {"B": 30, "C": 24, "D": 20, **{get_column_letter(c): 12 for c in range(5, 12)}})

    c = merge(ws, "B5:D5")
    c.value = "P(0 лідів) = e^(−Спенд ÷ Max CPL)"
    c.font, c.fill = font(TXT, 14, True, name=MONO), fill(PANEL)
    c.alignment = Alignment(vertical="center", indent=1)
    c = merge(ws, "B6:D8")
    c.value = ("Ймовірність, що адсет, який насправді в цілі, не дав жодного ліда при такому спенді. "
               "У Excel: =EXP(-Спенд/MaxCPL). Max CPL кожного гео береться з «Гео-калькулятор».")
    c.font, c.fill = font(MUTED, 9), fill(PANEL)
    c.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
    for rr in range(5, 9):
        for cc in range(2, 5):
            ws.cell(rr, cc).fill = fill(PANEL)
    ws.row_dimensions[5].height = 30

    label(ws["F4"], "ПОРОГИ ВЕРДИКТУ (можна змінювати)", ACCENT, 8, True)
    for rr, (txt, val) in enumerate([(f"«{V_EARLY}», якщо шанс ≥", 0.30),
                                     (f"«{V_Q}», якщо шанс ≥", 0.10),
                                     (f"«{V_OFF}», якщо шанс ≥", 0.03)], start=5):
        label(merge(ws, f"F{rr}:G{rr}"), txt, TXT, 9)
        ws[f"H{rr}"] = val
        inp(ws[f"H{rr}"], PCT0, "center")
    label(merge(ws, "F8:H8"), f"Нижче — «{V_LATE}»", MUTED, 9)

    group(ws, 10, 2, 4, "ТАБЛИЦЯ РІШЕНЬ")
    group(ws, 10, 5, 11, "СКІЛЬКИ ЦЕ В $ ДЛЯ ВАШИХ ГЕО (множник × Max CPL гео)", GREEN)
    headers(ws, 11, 2, ["Витрачено без лідів", "Шанс, що адсет насправді нормальний", "Що робити"])
    for i in range(N_GEO):
        col = 5 + i
        c = ws.cell(11, col, f'=IF({ref(S_CALC, f"$B${GEO_R0 + i}")}="","—",{ref(S_CALC, f"$B${GEO_R0 + i}")})')
        c.fill, c.font = fill(HEAD), font(HEAD_TXT, 9, True)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = Border(bottom=side(ACCENT, "medium"))
    ws.row_dimensions[11].height = 36

    for k, m in enumerate([0.5, 1, 1.5, 2, 2.5, 3, 4]):
        r = 12 + k
        ws[f"B{r}"] = m
        inp(ws[f"B{r}"], MULT, "left")
        ws[f"C{r}"] = f"=EXP(-B{r})"
        calc(ws[f"C{r}"], PCT0, bold=True)
        ws[f"D{r}"] = "=" + ladder(f"C{r}")
        calc(ws[f"D{r}"], None, "center", bold=True)
        for i in range(N_GEO):
            col = get_column_letter(5 + i)
            ws[f"{col}{r}"] = f'=IFERROR($B{r}*{ref(S_CALC, f"$K${GEO_R0 + i}")},"")'
            calc(ws[f"{col}{r}"], USD)
        ws.row_dimensions[r].height = 22
    cf_values(ws, "D12:D18", "$D12", VERDICT_STYLE)
    label(ws["B19"], "Множники в синіх клітинках можна змінити — все перерахується.", DIM, 8)

    # ── перевірка конкретних адсетів ──
    group(ws, 21, 2, 11, "ПЕРЕВІРКА АДСЕТІВ — вкажіть гео, спенд і кількість лідів")
    c = merge(ws, "B22:K23")
    c.value = ("Працює і коли ліди вже є: рахується шанс, що нормальний адсет дав би стільки або менше лідів "
               "(розподіл Пуассона, λ = Спенд ÷ Max CPL). Для 0 лідів це та сама формула e^(−λ). "
               f"Якщо CPL вже ≤ Max CPL — вердикт «{V_OK}».")
    c.font, c.alignment = font(MUTED, 9), Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[22].height = 18
    headers(ws, 24, 2, ["Адсет / назва", "Гео", "Спенд, $", "Лідів", "Max CPL, $",
                        "Очікувано лідів (λ)", "CPL факт, $", "Шанс, що нормальний", "Вердикт"])
    merge(ws, "J24:K24")
    for k in range(CHK_R1 - CHK_R0 + 1):
        r = CHK_R0 + k
        row = checks[k] if k < len(checks) else (None,) * 4
        for j, (val, fmt, al) in enumerate(zip(row, [None, None, USD_IN, INT_IN],
                                                ["left", "center", "right", "right"])):
            inp(ws.cell(r, 2 + j, val), fmt, al)
        ws[f"F{r}"] = f'=IF($C{r}="","",IFERROR({lookup("K", f"$C{r}")},""))'
        ws[f"G{r}"] = f'=IF(OR($D{r}="",F{r}=""),"",$D{r}/F{r})'
        ws[f"H{r}"] = f'=IF($D{r}="","",IFERROR($D{r}/$E{r},"—"))'
        ws[f"I{r}"] = f'=IF(G{r}="","",POISSON(N($E{r}),G{r},TRUE))'
        ws[f"J{r}"] = (f'=IF(I{r}="","",IF(AND(N($E{r})>0,$D{r}/MAX(N($E{r}),1)<=F{r}),"{V_OK}",'
                       + ladder(f"I{r}") + "))")
        for col, fmt in [("F", USD), ("G", "0.0"), ("H", USD), ("I", PCT0)]:
            calc(ws[f"{col}{r}"], fmt)
        merge(ws, f"J{r}:K{r}")
        calc(ws[f"J{r}"], None, "center", bold=True)
        ws.row_dimensions[r].height = 20
    cf_values(ws, f"J{CHK_R0}:K{CHK_R1}", f"$J{CHK_R0}", VERDICT_STYLE)
    dv_list(ws, f"C{CHK_R0}:C{CHK_R1}", GEO_LIST)


def build_stat(ws, rows):
    paint(ws, STAT_R1 + 4)
    title(ws, "Щоденна статистика",
          "1 рядок = день × гео (можна × кампанія). Апруви вносьте, коли прийдуть від ПП — "
          "до того дохід рахується прогнозно за планом апруву з калькулятора.")
    widths(ws, {"B": 11, "C": 7, "D": 24, "E": 11, "F": 11, "G": 9, "H": 8, "I": 10, "J": 8, "K": 8,
                "L": 9, "M": 8, "N": 9, "O": 9, "P": 9, "Q": 9, "R": 9, "S": 10, "T": 11, "U": 11,
                "V": 9, "W": 10})
    group(ws, 5, 2, 9, "ВВОДИТЕ ВИ")
    group(ws, 5, 10, 16, "ТРАФІК — АВТО", GREEN)
    group(ws, 5, 17, 23, "ГРОШІ — АВТО", GREEN)
    headers(ws, 6, 2, ["Дата", "Гео", "Кампанія / адсет", "Спенд, $", "Покази", "Кліки", "Ліди",
                       "Апруви (факт)", "CTR", "CPC, $", "CPM, $", "CR", "CPL, $", "Max CPL, $",
                       "CPL / Max", "Апруви*", "% апруву", "Ціна апруву, $", "Дохід, $", "Профіт, $",
                       "ROI", "Дані"])
    note(ws["I6"], "Залиште порожнім, поки ПП не віддала апруви. Тоді Апруви* = Ліди × план % апруву.")
    note(ws["Q6"], "Факт, якщо внесені апруви, інакше прогноз: Ліди × % апруву (план) з калькулятора.")
    note(ws["P6"], "CPL ÷ Max CPL. ≤100% — в цілі (зелений), до 130% — жовтий, вище — червоний.")

    for k in range(STAT_R1 - STAT_R0 + 1):
        r = STAT_R0 + k
        row = rows[k] if k < len(rows) else (None,) * 8
        for j, (val, fmt, al) in enumerate(zip(
                row, [DATE, None, None, USD_IN, INT_IN, INT_IN, INT_IN, INT_IN],
                ["center", "center", "left", "right", "right", "right", "right", "right"])):
            inp(ws.cell(r, 2 + j, val), fmt, al)
        on = f'$E{r}=""'
        f = {
            "J": (f'=IF({on},"",IFERROR(G{r}/F{r},""))', PCT1),
            "K": (f'=IF({on},"",IFERROR(E{r}/G{r},""))', USD),
            "L": (f'=IF({on},"",IFERROR(E{r}/F{r}*1000,""))', USD),
            "M": (f'=IF({on},"",IFERROR(H{r}/G{r},""))', PCT1),
            "N": (f'=IF({on},"",IFERROR(E{r}/H{r},"—"))', USD),
            "O": (f'=IF({on},"",IFERROR({lookup("K", f"$C{r}")},""))', USD),
            "P": (f'=IF({on},"",IFERROR(N{r}/O{r},""))', PCT0),
            "Q": (f'=IF({on},"",IF($I{r}<>"",$I{r},N($H{r})*IFERROR({lookup("E", f"$C{r}")},0)))', DEC1),
            "R": (f'=IF({on},"",IFERROR(Q{r}/H{r},""))', PCT0),
            "S": (f'=IF({on},"",IFERROR(E{r}/Q{r},"—"))', USD),
            "T": (f'=IF({on},"",Q{r}*IFERROR({lookup("D", f"$C{r}")},0))', USD),
            "U": (f'=IF({on},"",T{r}-E{r})', PROFIT),
            "V": (f'=IF({on},"",IFERROR(U{r}/E{r},""))', ROI),
            "W": (f'=IF({on},"",IF($I{r}<>"","факт","прогноз"))', None),
        }
        for col, (formula, fmt) in f.items():
            ws[f"{col}{r}"] = formula
            calc(ws[f"{col}{r}"], fmt, "center" if col == "W" else "right")
        ws.row_dimensions[r].height = 18

    rng = lambda col: f"{col}{STAT_R0}:{col}{STAT_R1}"
    cf_cpl_ratio(ws, rng("P"), f"$P{STAT_R0}")
    cf_sign(ws, f"U{STAT_R0}:V{STAT_R1}", f"U{STAT_R0}")
    cf(ws, rng("W"), f'$W{STAT_R0}="факт"', GREEN, bold=False)
    cf(ws, rng("W"), f'$W{STAT_R0}="прогноз"', MUTED, bold=False)
    dv_list(ws, rng("C"), GEO_LIST)
    ws.freeze_panes = f"E{STAT_R0}"


def build_day(ws):
    paint(ws, DAY_R1 + 30)
    title(ws, "Загальна статистика по днях",
          "Всі гео разом, рахується зі «Статистики». За замовчуванням — місяць останнього дня з даними.")
    widths(ws, {"B": 11, "C": 6, "D": 11, "E": 11, "F": 9, "G": 8, "H": 9, "I": 11, "J": 11,
                "K": 9, "L": 8, "M": 8, "N": 9, "O": 8, "P": 9, "Q": 10, "R": 13})
    sum_rng = lambda col: f"{col}{DAY_R0}:{col}{DAY_R1}"
    tiles = [
        ("СПЕНД", f"=SUM({sum_rng('D')})", USD0, TXT),
        ("ДОХІД", f"=SUM({sum_rng('I')})", USD0, TXT),
        ("ПРОФІТ", f"=SUM({sum_rng('J')})", PROFIT0, None),
        ("ROI", f'=IFERROR(SUM({sum_rng("J")})/SUM({sum_rng("D")}),"–")', ROI, None),
        ("ЛІДИ", f"=SUM({sum_rng('G')})", INT, TXT),
        ("CPL", f'=IFERROR(SUM({sum_rng("D")})/SUM({sum_rng("G")}),"–")', USD, TXT),
        ("АПРУВИ*", f"=SUM({sum_rng('H')})", INT, TXT),
        ("ЦІНА АПРУВУ", f'=IFERROR(SUM({sum_rng("D")})/SUM({sum_rng("H")}),"–")', USD, TXT),
    ]
    gap = side(BG, "thick")
    for k, (name, formula, fmt, color) in enumerate(tiles):
        c1 = 2 + 2 * k
        a, b = get_column_letter(c1), get_column_letter(c1 + 1)
        lab = merge(ws, f"{a}5:{b}5")
        val = merge(ws, f"{a}6:{b}6")
        lab.value, val.value = name, formula
        lab.font, val.font = font(MUTED, 8, True), font(color or TXT, 15, True)
        val.number_format = fmt
        for rr in (5, 6):
            for cc in (c1, c1 + 1):
                cell = ws.cell(rr, cc)
                cell.fill = fill(PANEL)
                cell.border = Border(left=gap if cc == c1 else None, right=gap if cc == c1 + 1 else None,
                                     top=side(ACCENT, "medium") if rr == 5 else None)
        lab.alignment = Alignment(horizontal="left", vertical="bottom", indent=1)
        val.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        if color is None:
            cf_sign(ws, f"{a}6", f"{a}6")
    ws.row_dimensions[5].height, ws.row_dimensions[6].height = 18, 30
    ws.row_dimensions[4].height = ws.row_dimensions[7].height = ws.row_dimensions[9].height = 6

    label(merge(ws, "B8:C8"), "Початок періоду", TXT, 9, True)
    inp(ws["D8"], DATE, "center")
    note(ws["D8"], "Порожньо = з 1-го числа місяця, в якому останній день у «Статистика» (показує весь місяць). "
                   "Вкажете дату — покаже 31 день від неї.")
    label(merge(ws, "E8:J8"), "порожньо = поточний місяць з даних; дата = 31 день від неї", DIM, 8)
    label(ws["L8"], "Період:", MUTED, 9, align="right")
    dates = stat_col("B")
    start = merge(ws, "M8:N8")
    start.value = (f'=IF($D$8<>"",$D$8,IF(COUNT({dates})=0,DATE(YEAR(TODAY()),MONTH(TODAY()),1),'
                   f'DATE(YEAR(MAX({dates})),MONTH(MAX({dates})),1)))')
    label(ws["O8"], "—", MUTED, 9, align="center")
    end = merge(ws, "P8:Q8")
    end.value = '=IF($D$8<>"",$D$8+30,EOMONTH($M$8,0))'
    for c in (start, end):
        c.number_format, c.font = DATE, font(TXT, 9, True)
        c.alignment = Alignment(horizontal="center", vertical="center")

    headers(ws, 10, 2, ["Дата", "День", "Спенд, $", "Покази", "Кліки", "Ліди", "Апруви*", "Дохід, $",
                        "Профіт, $", "ROI", "CTR", "CPC, $", "CPM, $", "CR", "CPL, $", "Ціна апруву, $",
                        "Профіт наростаючим, $"], 34)
    note(ws["H10"], "Факт там, де внесені апруви, і прогноз за планом апруву там, де ще ні.")

    def sumifs(col, r):
        return f"SUMIFS({stat_col(col)},{dates},$B{r})"

    for k in range(DAY_N):
        r = DAY_R0 + k
        ws[f"B{r}"] = "=$M$8" if k == 0 else f'=IF(B{r-1}="","",IF(B{r-1}+1>$P$8,"",B{r-1}+1))'
        on = f'$B{r}=""'
        f = {
            "C": (f'=IF({on},"",CHOOSE(WEEKDAY(B{r},2),"Пн","Вт","Ср","Чт","Пт","Сб","Нд"))', None),
            "D": (f'=IF({on},"",{sumifs("E", r)})', USD),
            "E": (f'=IF({on},"",{sumifs("F", r)})', INT),
            "F": (f'=IF({on},"",{sumifs("G", r)})', INT),
            "G": (f'=IF({on},"",{sumifs("H", r)})', INT),
            "H": (f'=IF({on},"",{sumifs("Q", r)})', DEC1),
            "I": (f'=IF({on},"",{sumifs("T", r)})', USD),
            "J": (f'=IF({on},"",I{r}-D{r})', PROFIT),
            "K": (f'=IF({on},"",IFERROR(J{r}/D{r},""))', ROI),
            "L": (f'=IF({on},"",IFERROR(F{r}/E{r},""))', PCT1),
            "M": (f'=IF({on},"",IFERROR(D{r}/F{r},""))', USD),
            "N": (f'=IF({on},"",IFERROR(D{r}/E{r}*1000,""))', USD),
            "O": (f'=IF({on},"",IFERROR(G{r}/F{r},""))', PCT1),
            "P": (f'=IF({on},"",IFERROR(D{r}/G{r},""))', USD),
            "Q": (f'=IF({on},"",IFERROR(D{r}/H{r},""))', USD),
            "R": (f'=IF(OR({on},$B{r}>MAX({dates})),"",SUM($J${DAY_R0}:J{r}))', PROFIT),
        }
        ws[f"B{r}"].number_format = DATE
        calc(ws[f"B{r}"], DATE, "center")
        for col, (formula, fmt) in f.items():
            ws[f"{col}{r}"] = formula
            calc(ws[f"{col}{r}"], fmt, "center" if col == "C" else "right", bold=(col == "J"))
        ws.row_dimensions[r].height = 18

    body = f"B{DAY_R0}:R{DAY_R1}"
    cf(ws, body, f'AND($B{DAY_R0}<>"",$B{DAY_R0}>MAX({dates}))', DIM, bold=False)
    cf(ws, f"C{DAY_R0}:C{DAY_R1}", f'OR($C{DAY_R0}="Сб",$C{DAY_R0}="Нд")', ACCENT, bold=False)
    cf_sign(ws, f"J{DAY_R0}:K{DAY_R1}", f"J{DAY_R0}")
    cf_sign(ws, f"R{DAY_R0}:R{DAY_R1}", f"R{DAY_R0}")

    # графік
    r_chart = DAY_R1 + 3
    label(ws[f"B{r_chart}"], "СПЕНД vs ДОХІД ПО ДНЯХ + ПРОФІТ НАРОСТАЮЧИМ", ACCENT, 9, True)
    bar = BarChart()
    bar.type, bar.grouping, bar.gapWidth, bar.overlap = "col", "clustered", 60, -10
    bar.add_data(Reference(ws, min_col=4, min_row=10, max_row=DAY_R1), titles_from_data=True)
    bar.add_data(Reference(ws, min_col=9, min_row=10, max_row=DAY_R1), titles_from_data=True)
    bar.set_categories(Reference(ws, min_col=2, min_row=DAY_R0, max_row=DAY_R1))
    bar_colors(bar, [BAR1, "2DD4BF"])
    bar.x_axis.number_format = "dd.mm"
    bar.y_axis.number_format = "$#,##0"
    line = LineChart()
    line.add_data(Reference(ws, min_col=18, min_row=10, max_row=DAY_R1), titles_from_data=True)
    s = line.series[0]
    s.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=YELLOW, w=28575))
    s.marker.symbol, s.smooth = "none", False
    line.y_axis.axId, line.y_axis.crosses = 200, "max"
    line.y_axis.majorGridlines = None
    line.y_axis.number_format = "$#,##0"
    dark_chart(bar, [bar.x_axis, bar.y_axis])
    dark_chart(line, [line.y_axis])
    bar += line
    bar.width, bar.height = 31, 9
    ws.add_chart(bar, f"B{r_chart + 1}")
    ws.freeze_panes = f"D{DAY_R0}"


def build_geo(ws):
    paint(ws, 45)
    title(ws, "Зведення по гео", "Підсумок по кожному гео за обраний період. План vs факт по апруву, статус гео.")
    widths(ws, {"B": 8, "C": 24, "D": 11, "E": 11, "F": 9, "G": 8, "H": 9, "I": 9, "J": 9, "K": 9,
                "L": 10, "M": 10, "N": 11, "O": 11, "P": 9, "Q": 9, "R": 10, "S": 10, "T": 18})
    label(ws["B5"], "Від", TXT, 9, True, "right")
    inp(ws["C5"], DATE, "center")
    label(ws["D5"], "До", TXT, 9, True, "right")
    merge(ws, "E5:F5")
    inp(ws["E5"], DATE, "center")
    label(merge(ws, "G5:O5"), "порожньо = весь час. Напр. вкажіть останні 7 днів, щоб бачити свіжу картину.", DIM, 8)
    ws.row_dimensions[5].height = 20

    dates = stat_col("B")
    period = f'{dates},">="&IF($C$5="",0,$C$5),{dates},"<="&IF($E$5="",2958465,$E$5)'
    group(ws, 6, 2, 3, "ГЕО")
    group(ws, 6, 4, 10, "ТРАФІК", GREEN)
    group(ws, 6, 11, 17, "ГРОШІ", GREEN)
    group(ws, 6, 18, 20, "АПРУВ І СТАТУС", YELLOW)
    headers(ws, 7, 2, ["Гео", "Оффер", "Спенд, $", "Покази", "Кліки", "Ліди", "CPL, $", "Max CPL, $",
                       "CPL / Max", "Апруви*", "Ціна апруву, $", "Max ціна апруву, $", "Дохід, $",
                       "Профіт, $", "ROI", "Ціль ROI", "% апруву факт", "% апруву план", "Статус"], 40)
    note(ws["R7"], "Апруви ÷ ліди тільки по днях, де апруви вже внесені. Якщо факт стабільно нижче плану — "
                   "оновіть план у калькуляторі, бо Max CPL зараз завищений.")

    for i in range(N_GEO):
        r, cr = 8 + i, GEO_R0 + i
        g = lambda col: ref(S_CALC, f"${col}${cr}")
        s = lambda col: f'SUMIFS({stat_col(col)},{stat_col("C")},$B{r},{period})'
        on = f'$B{r}=""'
        f = {
            "B": (f'=IF({g("B")}="","",{g("B")})', None, "center"),
            "C": (f'=IF({on},"",{g("C")})', None, "left"),
            "D": (f'=IF({on},"",{s("E")})', USD, None),
            "E": (f'=IF({on},"",{s("F")})', INT, None),
            "F": (f'=IF({on},"",{s("G")})', INT, None),
            "G": (f'=IF({on},"",{s("H")})', INT, None),
            "H": (f'=IF({on},"",IFERROR(D{r}/G{r},"—"))', USD, None),
            "I": (f'=IF({on},"",{g("K")})', USD, None),
            "J": (f'=IF({on},"",IFERROR(H{r}/I{r},""))', PCT0, None),
            "K": (f'=IF({on},"",{s("Q")})', DEC1, None),
            "L": (f'=IF({on},"",IFERROR(D{r}/K{r},"—"))', USD, None),
            "M": (f'=IF({on},"",{g("L")})', USD, None),
            "N": (f'=IF({on},"",{s("T")})', USD, None),
            "O": (f'=IF({on},"",N{r}-D{r})', PROFIT, None),
            "P": (f'=IF({on},"",IFERROR(O{r}/D{r},""))', ROI, None),
            "Q": (f'=IF({on},"",{g("F")})', PCT0, None),
            "R": (f'=IF({on},"",IFERROR({s("I")}/SUMIFS({stat_col("H")},{stat_col("C")},$B{r},'
                  f'{stat_col("I")},"<>",{period}),"—"))', PCT0, None),
            "S": (f'=IF({on},"",{g("E")})', PCT0, None),
            "T": (f'=IF({on},"",IF(N(D{r})=0,"Немає даних",IF(P{r}>=Q{r},"{V_SCALE}",'
                  f'IF(P{r}>=0,"Плюс, нижче цілі","Мінус"))))', None, "center"),
        }
        for col, (formula, fmt, al) in f.items():
            ws[f"{col}{r}"] = formula
            calc(ws[f"{col}{r}"], fmt, al or "right", bold=col in ("B", "O", "T"))
        ws.row_dimensions[r].height = 22

    r = 8 + N_GEO
    tot = {
        "B": ("Разом", None), "D": (f"=SUM(D8:D{r-1})", USD), "E": (f"=SUM(E8:E{r-1})", INT),
        "F": (f"=SUM(F8:F{r-1})", INT), "G": (f"=SUM(G8:G{r-1})", INT),
        "H": (f'=IFERROR(D{r}/G{r},"—")', USD), "K": (f"=SUM(K8:K{r-1})", DEC1),
        "L": (f'=IFERROR(D{r}/K{r},"—")', USD), "N": (f"=SUM(N8:N{r-1})", USD),
        "O": (f"=N{r}-D{r}", PROFIT), "P": (f'=IFERROR(O{r}/D{r},"")', ROI),
        "R": (f'=IFERROR(SUMIFS({stat_col("I")},{period})/SUMIFS({stat_col("H")},{stat_col("I")},"<>",{period}),"—")',
              PCT0),
    }
    for cc in range(2, 21):
        c = ws.cell(r, cc)
        c.fill, c.font = fill(HEAD), font(TXT, 10, True)
        c.border = Border(top=side(ACCENT, "medium"))
        c.alignment = Alignment(horizontal="right", vertical="center")
    for col, (val, fmt) in tot.items():
        ws[f"{col}{r}"] = val
        if fmt:
            ws[f"{col}{r}"].number_format = fmt
    ws[f"B{r}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    merge(ws, f"B{r}:C{r}")
    ws.row_dimensions[r].height = 24

    cf_cpl_ratio(ws, "J8:J14", "$J8")
    cf_sign(ws, f"O8:P{r}", f"O8")
    cf(ws, "R8:R14", "AND(ISNUMBER(R8),ISNUMBER(S8),R8<S8*0.9)", RED, bold=True)
    cf(ws, "R8:R14", "AND(ISNUMBER(R8),ISNUMBER(S8),R8>=S8)", GREEN, bold=True)
    cf_values(ws, "T8:T14", "$T8", [(V_SCALE, BG, GREEN), ("Плюс, нижче цілі", YELLOW, YELLOW_BG),
                                    ("Мінус", RED, RED_BG), ("Немає даних", DIM, None)])
    label(ws[f"B{r+1}"], "Апруви* — факт, де внесено, і прогноз за планом, де ще ні. "
                         "% апруву факт червоний, якщо на 10%+ нижче плану.", DIM, 8)

    rc = r + 3
    label(ws[f"B{rc}"], "СПЕНД vs ДОХІД ПО ГЕО", ACCENT, 9, True)
    bar = BarChart()
    bar.type, bar.grouping, bar.gapWidth = "col", "clustered", 80
    bar.add_data(Reference(ws, min_col=4, min_row=7, max_row=7 + N_GEO), titles_from_data=True)
    bar.add_data(Reference(ws, min_col=14, min_row=7, max_row=7 + N_GEO), titles_from_data=True)
    bar.set_categories(Reference(ws, min_col=2, min_row=8, max_row=7 + N_GEO))
    bar_colors(bar, [BAR1, "2DD4BF"])
    bar.y_axis.number_format = "$#,##0"
    dark_chart(bar, [bar.x_axis, bar.y_axis])
    bar.width, bar.height = 22, 8
    ws.add_chart(bar, f"B{rc + 1}")


def build_log(ws, actions):
    paint(ws, LOG_R1 + 4)
    title(ws, "Журнал дій",
          "Кожна зміна = 1 рядок. Метрики за N днів ДО і ПІСЛЯ дії підтягуються зі «Статистики» самі — "
          "видно, що до чого призвело.")
    widths(ws, {"B": 5, "C": 11, "D": 7, "E": 20, "F": 20, "G": 34, "H": 30, "I": 9, "J": 9, "K": 9,
                "L": 9, "M": 9, "N": 9, "O": 9, "P": 17, "Q": 36})
    label(merge(ws, "B5:C5"), "Вікно, днів", TXT, 9, True, "right")
    ws["D5"] = 3
    inp(ws["D5"], "0", "center")
    label(ws["E5"], "Поріг зміни CPL", TXT, 9, True, "right")
    ws["F5"] = 0.10
    inp(ws["F5"], PCT0, "center")
    c = merge(ws, "G5:Q5")
    label(c, "ДО = N днів перед датою дії, ПІСЛЯ = N днів після (день дії не рахується). Дані — по гео "
             "(і по кампанії, якщо вказана і так само записана в «Статистиці»). Порожнє гео = всі гео.", DIM, 8,
          wrap=True)
    ws.row_dimensions[5].height = 26

    group(ws, 7, 2, 8, "ЩО ЗРОБИЛИ — ЗАПОВНЮЄТЕ ВИ")
    group(ws, 7, 9, 16, "ЩО ЦЕ ДАЛО — АВТОМАТИЧНО", GREEN)
    group(ws, 7, 17, 17, "ВИСНОВОК", YELLOW)
    headers(ws, 8, 2, ["№", "Дата", "Гео", "Кампанія / адсет", "Тип дії", "Що зроблено", "Гіпотеза — чому",
                       "CPL до, $", "CPL після, $", "Δ CPL", "ROI до", "ROI після", "Лідів/день до",
                       "Лідів/день після", "Авто-оцінка", "Висновок / наступний крок"], 36)
    # службові колонки S:X (згорнуті)
    for j, h in enumerate(["Спенд до", "Ліди до", "Дохід до", "Спенд після", "Ліди після", "Дохід після"]):
        c = ws.cell(8, 19 + j, h)
        c.font, c.fill = font(MUTED, 8, True), fill(HEAD)
    ws.column_dimensions.group("S", "X", hidden=True)

    dates = stat_col("B")
    for k in range(LOG_R1 - LOG_R0 + 1):
        r = LOG_R0 + k
        a = actions[k] if k < len(actions) else (None,) * 7
        ws[f"B{r}"] = f'=IF($C{r}="","",ROW()-{LOG_R0 - 1})'
        calc(ws[f"B{r}"], "0", "center", color=MUTED)
        for j, (val, fmt, al) in enumerate(zip(a[:6], [DATE, None, None, None, None, None],
                                                ["center", "center", "left", "left", "left", "left"])):
            c = ws.cell(r, 3 + j, val)
            inp(c, fmt, al)
            if j >= 4:
                c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        q = ws.cell(r, 17, a[6])
        inp(q, None, "left")
        q.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        crit = (f'{stat_col("C")},IF($D{r}="","<>§",$D{r}),{stat_col("D")},IF($E{r}="","<>§",$E{r})')
        before = f'{dates},">="&($C{r}-$D$5),{dates},"<="&($C{r}-1),{crit}'
        after = f'{dates},">="&($C{r}+1),{dates},"<="&($C{r}+$D$5),{crit}'
        for j, (col, win) in enumerate([("E", before), ("H", before), ("T", before),
                                        ("E", after), ("H", after), ("T", after)]):
            c = ws.cell(r, 19 + j, f'=IF($C{r}="","",SUMIFS({stat_col(col)},{win}))')
            c.font, c.fill = font(MUTED, 8), fill(BG)
        on = f'$C{r}=""'
        f = {
            "I": (f'=IF({on},"",IFERROR(S{r}/T{r},"—"))', USD),
            "J": (f'=IF({on},"",IFERROR(V{r}/W{r},"—"))', USD),
            "K": (f'=IF({on},"",IFERROR(J{r}/I{r}-1,""))', '+0%;-0%;0%'),
            "L": (f'=IF({on},"",IFERROR((U{r}-S{r})/S{r},"—"))', ROI),
            "M": (f'=IF({on},"",IFERROR((X{r}-V{r})/V{r},"—"))', ROI),
            "N": (f'=IF({on},"",T{r}/$D$5)', DEC1),
            "O": (f'=IF({on},"",W{r}/$D$5)', DEC1),
            "P": (f'=IF({on},"",IF(MAX({dates})<$C{r}+$D$5,"Чекаємо даних",IF(S{r}=0,"Новий (без бази)",'
                  f'IF(OR(T{r}=0,W{r}=0),"Мало даних",IF(K{r}<=-$F$5,"Покращило",'
                  f'IF(K{r}>=$F$5,"Погіршило","Без змін"))))))', None),
        }
        for col, (formula, fmt) in f.items():
            ws[f"{col}{r}"] = formula
            calc(ws[f"{col}{r}"], fmt, "center" if col == "P" else "right", bold=(col == "P"))
        ws.row_dimensions[r].height = 30

    rng = lambda col: f"{col}{LOG_R0}:{col}{LOG_R1}"
    cf(ws, rng("K"), f"AND(ISNUMBER(K{LOG_R0}),K{LOG_R0}<=-$F$5)", GREEN)
    cf(ws, rng("K"), f"AND(ISNUMBER(K{LOG_R0}),K{LOG_R0}>=$F$5)", RED)
    cf_sign(ws, f"L{LOG_R0}:M{LOG_R1}", f"L{LOG_R0}")
    cf_values(ws, rng("P"), f"$P{LOG_R0}", [
        ("Покращило", GREEN, GREEN_BG), ("Погіршило", RED, RED_BG), ("Без змін", GRAY, GRAY_BG),
        ("Чекаємо даних", YELLOW, YELLOW_BG), ("Мало даних", MUTED, None), ("Новий (без бази)", ACCENT, None)])
    dv_list(ws, rng("D"), GEO_LIST)
    dv_list(ws, rng("F"), ref(S_LIST, "$B$4:$B$25"))
    ws.freeze_panes = f"D{LOG_R0}"


def build_creo(ws, creos):
    paint(ws, CRE_R1 + 4)
    title(ws, "Тест креативів", "Кожен креатив — рядок. Вердикт рахується так само, як у «Спенд без лідів», "
                                "плюс сигнал на масштабування.")
    widths(ws, {"B": 9, "C": 11, "D": 7, "E": 18, "F": 12, "G": 30, "H": 10, "I": 10, "J": 8, "K": 7,
                "L": 8, "M": 8, "N": 8, "O": 9, "P": 9, "Q": 9, "R": 10, "S": 17, "T": 30})
    label(merge(ws, "B5:D5"), "Масштабувати, якщо CPL ≤", TXT, 9, True, "right")
    ws["E5"] = 0.8
    inp(ws["E5"], PCT0, "center")
    label(merge(ws, "F5:G5"), "від Max CPL  і лідів ≥", TXT, 9, True)
    ws["H5"] = 5
    inp(ws["H5"], "0", "center")
    ws.row_dimensions[5].height = 20

    group(ws, 6, 2, 11, "КРЕАТИВ І ЦИФРИ — ЗАПОВНЮЄТЕ ВИ")
    group(ws, 6, 12, 19, "АНАЛІЗ — АВТОМАТИЧНО", GREEN)
    group(ws, 6, 20, 20, "")
    headers(ws, 7, 2, ["ID", "Запуск", "Гео", "Підхід / кут", "Формат", "Хук / опис", "Спенд, $", "Покази",
                       "Кліки", "Ліди", "CTR", "CPC, $", "CPM, $", "CPL, $", "Max CPL, $", "CPL / Max",
                       "Шанс, що норм.", "Вердикт", "Нотатки"], 36)
    for k in range(CRE_R1 - CRE_R0 + 1):
        r = CRE_R0 + k
        row = creos[k] if k < len(creos) else (None,) * 10
        for j, (val, fmt, al) in enumerate(zip(
                row, [None, DATE, None, None, None, None, USD_IN, INT_IN, INT_IN, INT_IN],
                ["left", "center", "center", "left", "center", "left", "right", "right", "right", "right"])):
            inp(ws.cell(r, 2 + j, val), fmt, al)
        inp(ws.cell(r, 20), None, "left")
        on = f'$H{r}=""'
        f = {
            "L": (f'=IF({on},"",IFERROR(J{r}/I{r},""))', PCT1),
            "M": (f'=IF({on},"",IFERROR(H{r}/J{r},""))', USD),
            "N": (f'=IF({on},"",IFERROR(H{r}/I{r}*1000,""))', USD),
            "O": (f'=IF({on},"",IFERROR(H{r}/K{r},"—"))', USD),
            "P": (f'=IF({on},"",IFERROR({lookup("K", f"$D{r}")},""))', USD),
            "Q": (f'=IF({on},"",IFERROR(O{r}/P{r},""))', PCT0),
            "R": (f'=IF(OR({on},P{r}=""),"",POISSON(N(K{r}),H{r}/P{r},TRUE))', PCT0),
            "S": (f'=IF(R{r}="","",IF(AND(N(K{r})>=$H$5,H{r}/MAX(N(K{r}),1)<=P{r}*$E$5),"{V_SCALE}",'
                  f'IF(AND(N(K{r})>0,H{r}/MAX(N(K{r}),1)<=P{r}),"{V_OK}",{ladder(f"R{r}")})))', None),
        }
        for col, (formula, fmt) in f.items():
            ws[f"{col}{r}"] = formula
            calc(ws[f"{col}{r}"], fmt, "center" if col == "S" else "right", bold=(col == "S"))
        ws.row_dimensions[r].height = 18
    cf_cpl_ratio(ws, f"Q{CRE_R0}:Q{CRE_R1}", f"$Q{CRE_R0}")
    cf_values(ws, f"S{CRE_R0}:S{CRE_R1}", f"$S{CRE_R0}", VERDICT_STYLE)
    dv_list(ws, f"D{CRE_R0}:D{CRE_R1}", GEO_LIST)
    dv_list(ws, f"F{CRE_R0}:F{CRE_R1}", ref(S_LIST, "$D$4:$D$15"))
    ws.freeze_panes = f"C{CRE_R0}"


ACTION_TYPES = ["Запуск", "Вимкнення", "Збільшення бюджету", "Зменшення бюджету", "Новий креатив",
                "Новий преленд / ленд", "Зміна таргету / аудиторії", "Зміна ставки / стратегії", "Дублювання",
                "Новий акаунт / BM", "Зміна оффера / ПП", "Інше"]
FORMATS = ["Відео", "UGC", "Статика", "Карусель", "Слайд-шоу", "Інше"]


def build_lists(ws):
    paint(ws, 30, 10)
    title(ws, "Списки", "Значення для випадаючих списків. Можна дописувати свої.")
    widths(ws, {"B": 28, "C": 3, "D": 20})
    headers(ws, 3, 2, ["Типи дій (Журнал дій)"], 22)
    headers(ws, 3, 4, ["Формати (Креативи)"], 22)
    for k in range(22):
        inp(ws.cell(4 + k, 2, ACTION_TYPES[k] if k < len(ACTION_TYPES) else None), None, "left")
    for k in range(12):
        inp(ws.cell(4 + k, 4, FORMATS[k] if k < len(FORMATS) else None), None, "left")


def build_start(ws, demo):
    paint(ws, 50, 16)
    widths(ws, {"B": 4, "C": 16, "D": 14, **{get_column_letter(c): 11 for c in range(5, 15)}})
    ws["B2"] = "Трекер арбітражу · Nutra / CPA"
    ws["B2"].font = font(TXT, 20, True)
    ws.row_dimensions[2].height = 32
    ws["B3"] = ("Ліміти по гео, щоденна статистика, журнал дій і рішення «вимикати чи чекати» — "
                "все рахується автоматично.")
    ws["B3"].font = font(MUTED, 10)

    def section(r, text):
        label(ws[f"B{r}"], text, ACCENT, 9, True)
        for cc in range(2, 15):
            ws.cell(r, cc).border = Border(bottom=side(BORDER))
        ws.row_dimensions[r].height = 22

    section(5, "ЯК КОРИСТУВАТИСЯ")
    steps = [
        "«Гео-калькулятор» — по кожному гео внесіть виплату за апрув, план % апруву, цільовий ROI, CR і CTR. "
        "EPL, Max CPL, Max CPC, Max CPM і Max ціна апруву порахуються самі.",
        "«Статистика» — щодня 1 рядок на гео (або гео × кампанія): спенд, покази, кліки, ліди. Апруви — коли "
        "прийдуть від ПП; до того дохід рахується прогнозно за планом апруву.",
        "«Журнал дій» — кожну зміну (запуск, бюджет, креатив, преленд, вимкнення) одним рядком. CPL, ROI і ліди "
        "за N днів до та після підтягнуться самі + авто-оцінка «покращило / погіршило».",
        "«Загальна по днях» і «Зведення по гео» — заповнюються самі. Змінюйте «Початок періоду» або «Від–До», "
        "щоб дивитися інший місяць чи тиждень.",
        "«Спенд без лідів» — таблиця рішень зі скріну + перевірка конкретних адсетів (і коли ліди вже є).",
        "«Креативи» — трекер тестів креативів з автоматичним вердиктом «масштабувати / в цілі / вимикати».",
    ]
    for k, text in enumerate(steps):
        r = 6 + k
        c = ws[f"B{r}"]
        c.value, c.font = k + 1, font(BG, 10, True)
        c.fill, c.alignment = fill(ACCENT), Alignment(horizontal="center", vertical="center")
        label(merge(ws, f"C{r}:N{r}"), text, TXT, 10, wrap=True)
        ws[f"C{r}"].alignment = Alignment(wrap_text=True, vertical="center", indent=1)
        ws.row_dimensions[r].height = 32

    section(13, "ПОЗНАЧЕННЯ")
    legend = [
        ("28.00", "inp", "Сині клітинки — вводите ви."),
        ("11.20", "calc", "Білий текст на темному — формула. Не перезаписуйте."),
        (V_SCALE, (BG, GREEN), "Сильно краще цілі — можна масштабувати."),
        (V_OK, (GREEN, GREEN_BG), "CPL ≤ Max CPL — тримати."),
        (V_EARLY, (GRAY, GRAY_BG), "Даних ще замало для висновку."),
        (V_Q, (YELLOW, YELLOW_BG), "Дивіться уважно — рішення найближчим часом."),
        (V_OFF, (RED, RED_BG), "Статистично адсет / креатив не в цілі."),
        (V_LATE, ("FFFFFF", DRED_BG), "Мали вимкнути раніше — гроші вже злиті."),
    ]
    for k, (sample, kind, text) in enumerate(legend):
        r = 14 + k
        c = merge(ws, f"C{r}:D{r}")
        c.value = sample
        if kind == "inp":
            inp(c, None, "center")
        elif kind == "calc":
            calc(c, None, "center")
        else:
            c.font, c.fill = font(kind[0], 9, True), fill(kind[1])
            c.alignment = Alignment(horizontal="center", vertical="center")
        label(merge(ws, f"E{r}:N{r}"), text, MUTED, 9)
        ws[f"E{r}"].alignment = Alignment(vertical="center", indent=1)
        ws.row_dimensions[r].height = 20

    section(23, "ГОЛОВНІ ФОРМУЛИ")
    formulas = [
        ("EPL", "Виплата × % апруву", "Скільки в середньому заробляєте з 1 ліда."),
        ("Max CPL", "EPL ÷ (1 + цільовий ROI)", "Ліміт CPL для цільового ROI. Беззбиток = EPL."),
        ("Max ціна апруву", "Виплата ÷ (1 + цільовий ROI)", "Скільки можна витратити на 1 апрув."),
        ("Max CPC", "Max CPL × CR", "Максимальна ціна кліка."),
        ("Max CPM", "Max CPC × CTR × 1000", "Максимальна ціна 1000 показів."),
        ("Ціна апруву", "Спенд ÷ Апруви", "Скільки фактично коштує 1 апрув."),
        ("ROI", "(Дохід − Спенд) ÷ Спенд", "Окупність."),
        ("P(0 лідів)", "e^(−Спенд ÷ Max CPL)", "Шанс, що нормальний адсет не дав лідів."),
    ]
    for k, (name, formula, text) in enumerate(formulas):
        r = 24 + k
        label(merge(ws, f"C{r}:D{r}"), name, TXT, 10, True)
        c = merge(ws, f"E{r}:H{r}")
        c.value, c.font = formula, font(ACCENT, 10, name=MONO)
        label(merge(ws, f"I{r}:N{r}"), text, MUTED, 9)
        ws.row_dimensions[r].height = 19

    section(33, "АРКУШІ")
    sheets = [
        (S_CALC, "налаштування гео й ліміти"), (S_STAT, "щоденне введення цифр"),
        (S_LOG, "що робили і що це дало"), (S_DAY, "всі гео разом по днях + графік"),
        (S_GEO, "підсумок і статус кожного гео"), (S_KILL, "вимикати чи чекати"),
        (S_CRE, "тести креативів"), (S_LIST, "значення випадаючих списків"),
    ]
    for k, (name, text) in enumerate(sheets):
        r = 34 + k
        c = merge(ws, f"C{r}:D{r}")
        c.value = name
        c.hyperlink = Hyperlink(ref=f"C{r}", location=f"'{name}'!A1", display=name)
        c.font = font(INPUT_TXT, 10, True, underline="single")
        label(merge(ws, f"E{r}:N{r}"), text, MUTED, 9)
        ws.row_dimensions[r].height = 19

    r = 43
    c = merge(ws, f"B{r}:N{r + 1}")
    if demo:
        c.value = ("ДЕМО: у файлі тестові дані за 01–12.10.2026, щоб було видно, як усе працює. Для роботи "
                   "беріть arbitrage_tracker.xlsx (чистий) або очистіть сині клітинки в «Статистика», "
                   "«Журнал дій», «Креативи» і в перевірці адсетів.")
    else:
        c.value = ("У «Статистика», «Журнал дій», «Креативи» і в перевірці адсетів — по одному рядку-прикладу. "
                   "Замініть їх своїми даними. Гео в калькуляторі — теж приклад, впишіть свої.")
    c.font, c.fill = font(YELLOW, 9, True), fill(YELLOW_BG)
    c.alignment = Alignment(wrap_text=True, vertical="center", indent=1)


# ═════════════════════════ ДАНІ ═════════════════════════
GEOS = [  # код, назва, виплата, апрув, ROI, CR, CTR, нотатка
    ("PL", "Польща · суглоби", 28, 0.40, 0.30, 0.050, 0.020, None),
    ("RO", "Румунія · діабет", 22, 0.45, 0.30, 0.055, 0.020, None),
    ("IT", "Італія · схуднення", 35, 0.35, 0.30, 0.050, 0.020, None),
    ("ES", "Іспанія · потенція", 33, 0.38, 0.30, 0.055, 0.018, None),
    ("CZ", "Чехія · тиск", 26, 0.40, 0.30, 0.050, 0.019, "тест з 06.10"),
    ("BG", "Болгарія · суглоби", 20, 0.42, 0.30, 0.060, 0.024, None),
    ("GR", "Греція · простатит", 27, 0.43, 0.30, 0.055, 0.021, None),
]
D = lambda day: dt.date(2026, 10, day)


def simulate():
    """Правдоподібна статистика за 01–12.10 з ефектами від дій у журналі."""
    rnd = random.Random(11)

    def jitter(x, s=0.10):
        return x * math.exp(rnd.gauss(0, s))

    def poisson(lam):
        k, p, lim = 0, 1.0, math.exp(-lam)
        while True:
            p *= rnd.random()
            if p <= lim:
                return k
            k += 1

    def params(geo, day):
        if geo == "PL":
            late = day >= D(8)
            return dict(camp="PL | суглоби | CBO-1", spend=225 if late else 150, cpm=9.4 if late else 7.5,
                        ctr=0.021, cr=0.05, appr=0.38)
        if geo == "RO":
            return dict(camp="RO | діабет | CBO-1", spend=175 if day >= D(9) else 120, cpm=6.5, ctr=0.02,
                        cr=0.055, appr=0.47)
        if geo == "IT":
            new = day >= D(6)
            return dict(camp="IT | схуднення | ABO", spend=170 if day >= D(12) else 130, cpm=11,
                        ctr=0.022 if new else 0.016, cr=0.05 if new else 0.04, appr=0.33)
        if geo == "ES":
            return dict(camp="ES | потенція | ABO", spend=110, cpm=9.5, ctr=0.017,
                        cr=0.06 if day >= D(5) else 0.042, appr=0.40)
        if geo == "CZ" and day >= D(6):
            return dict(camp="CZ | тиск | тест", spend=70, cpm=8, ctr=0.019, cr=0.05, appr=0.38)
        return None

    rows = []
    for d in range(1, 13):
        day = D(d)
        for geo in ["PL", "RO", "IT", "ES", "CZ"]:
            p = params(geo, day)
            if not p:
                continue
            spend = round(jitter(p["spend"], 0.06), 2)
            impr = round(spend / jitter(p["cpm"], 0.06) * 1000)
            clicks = round(impr * jitter(p["ctr"], 0.06))
            leads = poisson(clicks * p["cr"])
            appr = sum(rnd.random() < p["appr"] for _ in range(leads)) if day <= D(9) else None
            rows.append((day, geo, p["camp"], spend, impr, clicks, leads, appr))
    return rows


ACTIONS = [  # дата, гео, кампанія, тип, що, чому, висновок
    (D(4), "ES", None, "Новий преленд / ленд", "Замінив преленд на новий (історія лікаря)",
     "Старий преленд: CR 4%, CPL вище Max", "Преленд робочий — лишаємо, тестуємо його на IT"),
    (D(5), "IT", None, "Новий креатив", "3 UGC-відео замість статики",
     "Статика вигоріла, CTR 1.6%", "UGC заходить: CTR вище, CPL майже в цілі"),
    (D(6), "CZ", None, "Запуск", "Тестовий запуск, 2 адсети по $35", "Висока виплата, мало конкурентів", None),
    (D(7), "PL", None, "Збільшення бюджету", "Бюджет CBO $150 → $225",
     "6 днів стабільний ROI 50%+", "CPM виріс — бюджет назад до $180, масштаб через дублі"),
    (D(8), "RO", None, "Дублювання", "Дубль кращого адсету на 35–65, широка",
     "Масштаб без підняття бюджету основного", None),
    (D(11), "IT", None, "Збільшення бюджету", "+30% бюджету", "CPL в цілі 5 днів поспіль", None),
]

CREOS = [  # ID, запуск, гео, підхід, формат, хук, спенд, покази, кліки, ліди
    ("IT-01", D(3), "IT", "Біль / страх", "Статика", "«Лікар попереджає: ...»", 60, 5400, 85, 2),
    ("IT-02", D(5), "IT", "Історія", "UGC", "«Бабуся з Неаполя скинула 12 кг»", 140, 12000, 270, 16),
    ("IT-03", D(5), "IT", "До / після", "Відео", "Склейка до/після за 30 днів", 35, 3100, 41, 0),
    ("PL-04", D(2), "PL", "Експерт", "Відео", "Ортопед про біль у колінах", 210, 27000, 600, 33),
    ("RO-02", D(1), "RO", "Новина / ТВ", "Статика", "Скрін новинного сюжету", 95, 14500, 290, 15),
    ("ES-05", D(4), "ES", "Відгук", "UGC", "Відгук чоловіка 50+", 9, 950, 16, 0),
    ("ES-06", D(4), "ES", "Біль", "Слайд-шоу", "5 ознак, що пора до лікаря", 25, 2600, 40, 1),
    ("CZ-01", D(6), "CZ", "Експерт", "Відео", "Кардіолог про тиск зранку", 46, 5700, 110, 6),
]

CHECKS = [  # адсет, гео, спенд, лідів
    ("PL | 35-55 Ж | LAL 3%", "PL", 9, 0),
    ("RO | широка | дубль", "RO", 16, 0),
    ("IT | статика v2", "IT", 30, 0),
    ("ES | 45+ | інтереси", "ES", 40, 0),
    ("CZ | широка", "CZ", 25, 1),
    ("PL | відео ортопед", "PL", 60, 9),
]


def build(demo, path):
    wb = Workbook()
    ws = wb.active
    ws.title = S_START
    order = [S_CALC, S_STAT, S_LOG, S_DAY, S_GEO, S_KILL, S_CRE, S_LIST]
    sheets = {S_START: ws, **{n: wb.create_sheet(n) for n in order}}
    tabs = {S_START: YELLOW, S_CALC: ACCENT, S_STAT: "3B82F6", S_LOG: "3B82F6", S_DAY: GREEN,
            S_GEO: GREEN, S_KILL: RED, S_CRE: "3B82F6", S_LIST: DIM}
    for name, color in tabs.items():
        sheets[name].sheet_properties.tabColor = color

    stat = simulate() if demo else [(D(1), "PL", "PL | суглоби | CBO-1", 150, 20000, 420, 21, 8)]
    build_start(sheets[S_START], demo)
    build_calc(sheets[S_CALC], GEOS)
    build_stat(sheets[S_STAT], stat)
    build_log(sheets[S_LOG], ACTIONS if demo else ACTIONS[:1])
    build_day(sheets[S_DAY])
    build_geo(sheets[S_GEO])
    build_kill(sheets[S_KILL], CHECKS if demo else CHECKS[:1])
    build_creo(sheets[S_CRE], CREOS if demo else CREOS[:1])
    build_lists(sheets[S_LIST])
    if not demo:
        for sheet, cell in [(S_STAT, f"B{STAT_R0}"), (S_LOG, f"C{LOG_R0}"), (S_CRE, f"B{CRE_R0}"),
                            (S_KILL, f"B{CHK_R0}")]:
            note(sheets[sheet][cell], "Рядок-приклад — замініть своїми даними.")
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)


if __name__ == "__main__":
    out = Path(sys.argv[1] if len(sys.argv) > 1 else Path(__file__).parent)
    out.mkdir(parents=True, exist_ok=True)
    build(True, out / "arbitrage_tracker_demo.xlsx")
    build(False, out / "arbitrage_tracker.xlsx")
    print("ok:", out)
